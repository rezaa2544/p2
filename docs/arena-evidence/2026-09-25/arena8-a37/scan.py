#!/usr/bin/env python3
"""Read-only A-37 static inventory. No repository writes or certification claims."""
import pathlib,json,re,collections,csv,gzip,hashlib,posixpath
ROOT=pathlib.Path('/home/user/p2');OUT=pathlib.Path('/home/user/arena8-a37');TMP=pathlib.Path('/var/tmp/arena8-a37')
files=json.loads((OUT/'evidence/text-files.json').read_text());texts={f:(ROOT/f).read_text() for f in files};facts=json.loads((TMP/'evidence/ast.json').read_text());tests=sorted(f for f in facts if f.startswith('tests/'));findings=[];edges=[]
patterns={'assert(true)':r'\bassert\s*\(\s*true\b','process.exit(0)':r'\bprocess\s*\.\s*exit\s*\(\s*0\s*\)','0/0':r'(?<![\w\d])0\s*/\s*0(?!\d)','|| true':r'\|\|\s*true\b','MOCK':r'\b(?:mock\w*|stub\w*|fake\w*|pg-mem|jsdom)\b','hidden skip':r'\b(?:skip(?:ped|ping)?|NOT[- ]RUN|continue-on-error)\b','swallowed catch':r'\bcatch\s*(?:\([^)]*\))?\s*\{\s*(?:/\*[\s\S]*?\*/\s*)?\}'}
raw=collections.Counter();rawfiles=collections.defaultdict(set);rows_seen=set()
def zone(f,pos):
 v=facts.get(f,{})
 if f.endswith(('.sh','.yml','.yaml','.py')) and texts[f][:pos].split('\n')[-1].lstrip().startswith('#'):return 'comment'
 if any(a<=pos<b for a,b in v.get('comments',[])):return 'comment'
 if any(a<=pos<b for a,b in v.get('strings',[])):return 'string/template'
 if f.endswith(('.md','.json','.txt','.csv','.svg')):return 'documentation/data'
 return 'code-or-unparsed'
def snippet(f,line):return texts[f].splitlines()[line-1].strip()[:210] if texts[f].splitlines() else ''
def helper(f):return '/helpers/' in f or re.search(r'(?:-child|-worker|-browser|-lib|/runner|/index\.test|/master[^/]*|/mutant[^/]*)\.[cm]?js$',f)
def classify(f,kind,line,origin,detail):
 s=texts[f];lines=s.splitlines();around='\n'.join(lines[max(0,line-7):line+5]);ctx=around.lower()
 if f.endswith('.html') or f.startswith('reza/') or 'generated' in f or f=='authz/write-perms.json':return 'generated','INFO','static','deduplicate against source; generated copy is not an independent defect'
 if origin in ('comment','documentation/data'):return 'legitimate','INFO','static','nonexecuting text; do not count as behavior'
 if kind=='assert(true)' and origin.startswith('AST'):
  if f=='tests/report2.js':return 'real defect','P2','source-confirmed','cleanup success assertion is tautological; assert removed entities/relations'
  if f=='tests/simulation.js':return 'legitimate','INFO','manual-source-review','no-throw navigation smoke test; not proof of rendered correctness'
 if f=='tests/migration-009-live.js' and kind in ('process.exit(0)','0/0','hidden skip'):return 'real defect','P1','runtime-reproduced','missing dependency returns green 0/0; require nonzero NOT VERIFIED or enforce explicit required mode'
 if kind=='MOCK':
  if f=='tests/infrastructure/phase6/persistence.test.js':return 'certification blocker','P1','manual-source-review','Map-backed mock cannot establish PostgreSQL/container restart claims; keep unit label and add live evidence'
  if f.startswith(('tests/','tools/')):return 'fixture','INFO','static-candidate','mock/unit context is legitimate but cannot substitute for claimed live runtime evidence'
  return 'legitimate','INFO','static-candidate','mock/fake token alone is not proof of an active production mock; review context'
 if kind=='hidden skip' and f=='tests/smoke.js' and line==647:return 'real defect','P1','runtime-reproduced','missing fixture causes wrapper PASS with zero assertions; register NOT-RUN separately or fail required fixture'
 if kind=='hidden skip':
  if 'skip-link' in ctx or 'skipstr' in ctx or 'skipws' in ctx or origin=='comment':return 'legitimate','INFO','static','accessibility/parser/control-flow use, not a skipped test'
  if f.startswith('tests/api/') and 'no second manager' in ctx:return 'certification blocker','P1','manual-source-review','required cross-tenant negative test omitted when fixture absent; seed required manager and fail if absent'
  if re.search(r'process\.exit\(\s*[12]\s*\)',around) and 'exit(0)' not in around:return 'legitimate','INFO','static','missing prerequisite visibly fails closed; skip wording is not a green skip'
  if f=='tests/chat9-behavior-regression.js':return 'legitimate','P2','manual-source-review','optional NOT-RUN is separately counted; do not upgrade optional census to verified'
  if f.endswith('current-head-sync-pg.cjs'):return 'legitimate','INFO','manual-source-review','explicit diagnostic opt-out and scope disclosure; no silent full-browser claim'
  return 'certification blocker','P2','needs-review','skip/early return candidate: prove exclusion contract and executed-count propagation before clearance'
 if kind=='swallowed catch':
  if re.search(r'\b(kill|close|unlink|rmSync|rm\(|cleanup|disconnect|stop|removeListener)\b',around):return 'legitimate','P2','static-candidate','cleanup/best-effort context; verify primary failure remains observable'
  if f.startswith('tests/') and helper(f):return 'helper','P2','static-candidate','helper fallback; parent must assert outcome and propagate failure'
  if origin=='string/template':return 'fixture','P2','static-candidate','embedded code/template: may execute via eval/VM; not ordinary source execution proof'
  return 'certification blocker','P2','needs-review','error-to-success/fallback candidate; demonstrate error remains visible to required assertion/metric; do not blanket-remove catch'
 if kind=='process.exit(0)':
  if helper(f):return 'helper','INFO','static-candidate','worker/runner success exit; verify parent protocol/assertions'
  if f.startswith('tests/') and re.search(r'(?:fail\w*|bad|errors|failed)\b',around,re.I):return 'legitimate','INFO','static-candidate','guarded success footer candidate; zero-count and skip accounting still require validation'
  if not f.startswith('tests/'):return 'legitimate','INFO','static-candidate','CLI/normal shutdown success is not a test pass by itself'
  return 'certification blocker','P2','needs-review','success exit without locally evident failure/count guard; inspect complete control flow'
 if kind=='|| true':
  if f=='.github/workflows/observability-regression.yml':return 'legitimate','INFO','manual-source-review','polling is bounded and ends exit1 if required streams are absent'
  if origin.startswith('AST'):return 'legitimate','P2','needs-review','JavaScript fallback, not shell exit masking; inspect predicate semantics'
  if re.search(r'(?:pkill|kill |rm |down |stop |grep|git remote)',around):return 'legitimate','P2','static-candidate','cleanup or no-match capture candidate; ensure verdict is checked downstream'
  return 'certification blocker','P2','needs-review','shell status suppression; prove independent failing assertion or remove mask'
 if kind=='0/0':return 'legitimate','INFO','static-candidate','literal ratio/NaN/documentation token; no executed zero-check suite established from token alone'
 return 'legitimate','INFO','static-candidate','lexical candidate; not runtime evidence'
def add(f,kind,line,origin,detail,override=None):
 key=(f,kind,line,origin)
 if key in rows_seen:return
 rows_seen.add(key);cl,sev,confidence,fix=override or classify(f,kind,line,origin,detail)
 findings.append({'id':'','kind':kind,'file':f,'line':line,'origin':origin,'classification':cl,'severity':sev,'confidence':confidence,'detail':detail,'suggested_fix':fix,'snippet':snippet(f,line)})
for f,s in texts.items():
 for kind,pattern in patterns.items():
  for m in re.finditer(pattern,s,re.I if kind in ('MOCK','hidden skip') else 0):
   raw[kind]+=1;rawfiles[kind].add(f);line=s.count('\n',0,m.start())+1;z=zone(f,m.start())
   # JS AST is authoritative for executable constructs; retain lexical artifacts separately.
   if not (z!='comment' and any(a['kind']==kind and a['line']==line for a in facts.get(f,{}).get('markers',[]))):
    add(f,kind,line,z,'lexical marker; multiple occurrences on same line collapse in inventory')
 for m in facts.get(f,{}).get('markers',[]):add(f,m['kind'],m['line'],'AST',m['detail'])
# Static reference graph: explicit workflow paths + module dependencies + declared runner suites.
known=set(files)
def resolve(f,v):
 candidates=[v,posixpath.normpath(posixpath.join(posixpath.dirname(f),v))]
 for c in candidates:
  for x in [c,c+'.js',c+'.cjs',c+'/index.js']:
   if x in known:return x
 return None
for f,v in facts.items():
 s=texts[f]
 for ref in v.get('references',[]):
  val=ref['value'];dest=resolve(f,val)
  if not dest or dest==f:continue
  ln=snippet(f,ref['line']);typ='literal-reference-only'
  if re.search(r'\brequire\s*\(',ln):typ='module-dependency'
  elif re.search(r'(?:runner|index\.test|master|suite|mutations)',f) and re.search(r'\b(?:execSync|execFileSync|spawnSync|spawn)\b',s):typ='declared-runner-child'
  elif re.search(r'\b(?:execSync|execFileSync|spawnSync|spawn|fork)\b',ln):typ='literal-process-child'
  if typ!='literal-reference-only':edges.append((f,dest,ref['line'],typ))
# Generic test/tool/shell path invocation references in uncommented workflow bodies.
roots=set();direct=set();legacydirect=set()
for f,s in texts.items():
 if f.startswith('.github/workflows/'):
  legacydirect.update(re.findall(r'tests/[A-Za-z0-9_.\-/]+\.js',s))
  active='\n'.join(l for l in s.splitlines() if not l.lstrip().startswith('#'))
  for n,l in enumerate(s.splitlines(),1):
   if l.lstrip().startswith('#'):continue
   for dest in re.findall(r'(?:tests|tools|scripts)/[A-Za-z0-9_.\-/]+\.(?:[cm]?js|sh)',l):
    if dest in known:roots.add(dest);edges.append((f,dest,n,'workflow-reference-conditional'))
  direct.update(x for x in roots if x.startswith('tests/'))
pkg=json.loads(texts['package.json']);npmroots=set(re.findall(r'tests/[A-Za-z0-9_.\-/]+\.(?:[cm]?js)',pkg['scripts']['test']));roots.update(npmroots);direct.update(npmroots);legacydirect.update(npmroots)
def closure(start):
 seen=set(start);changed=True
 while changed:
  changed=False
  for a,b,_,_ in edges:
   if a in seen and b not in seen:seen.add(b);changed=True
 return seen
reachable=closure(roots);manual=set(f for f in tests if f.count('/')==1 and f.endswith('.js') and not f.endswith('-child.js'));manual.add('tests/api/runner.js');manualReach=closure(manual)
narrow=[];noNamed=[];noOracle=[];orphan=[]
for f in tests:
 s=texts[f];v=facts[f];top=f.count('/')==1;custom=bool(re.search(r'(?:function\s+(?:test|T|check|record)|(?:const|let|var)\s+(?:T|test|record)\s*=)',s) and re.search(r'fail\w*\s*(?:\+\+|\+=)|result\s*:\s*\w+\s*\?\s*[\'\"]PASS|\bFAIL\b',s));delegate=any(a==f and b.startswith('tests/') for a,b,_,_ in edges)
 if not re.search(r'\bassert\s*[.(]',s):narrow.append(f)
 if not v.get('checks'):noNamed.append(f)
 if not v.get('checks') and not v.get('throws') and not custom and not delegate and not helper(f):noOracle.append(f)
 if f in narrow or f in noNamed:
  if f in ['tests/a11y-interactive.js','tests/a11y-runtime.js','tests/demo-thursday.js']:ov=('legitimate','INFO','manual-source-review','runtime result/error/DOM guard is the oracle; absence of assert-named call is not zero checks')
  elif f=='tests/pg-outage-control.js':ov=('helper','INFO','manual-source-review','fault-control utility returns success/failure handles; caller owns assertion')
  elif helper(f) or delegate:ov=('helper','INFO','static','delegated/worker checks; execute parent and verify failure propagation')
  elif v.get('checks') or v.get('throws') or custom:ov=('legitimate','INFO','static','alternative assertion/throw/custom result mechanism exists; raw assert grep is insufficient')
  elif '-mutations' in f:ov=('helper','P2','needs-review','mutation harness may delegate checks; require killed/survived counts and child status')
  else:ov=('certification blocker','P1','needs-review','no recognized oracle; execute with counted assertions and negative control, reject zero checks')
  add(f,'ZERO-CHECK',1,'file-analysis',f'no_raw_assert={f in narrow}; no_named_check_calls={f in noNamed}; no_recognized_oracle={f in noOracle}',ov)
 if f not in reachable:
  orphan.append(f)
  if helper(f):ov=('helper','INFO','static-candidate','worker/helper may be intentionally indirect; record parent ownership')
  elif f in manualReach:ov=('certification blocker','P2','static-candidate','manual broad runner covers file but no proven CI path; wire required suite or document owner/schedule')
  elif '/performance/' in f:ov=('legitimate','P2','static-candidate','manual/load suite candidate; require separate scheduled runtime evidence for claims')
  else:ov=('certification blocker','P2','needs-review','no path in conservative CI graph; inspect dynamic dispatch before declaring definite orphan')
  add(f,'ORPHAN',1,'CI-graph','No static path from declared CI roots (includes dependency/runner overapproximation; not proof of never executed)',ov)
# Gate/design findings not represented by simple tokens.
for line,detail in [(26,'pattern-wide allowlist ignores file/line and owner/expiry governance'),(28,'empty registry iterates zero requirements and can return green')]:
 add('tools/strict-verification-gate.js','ZERO-CHECK' if line==28 else 'hidden skip',line,'manual-probe',detail,('real defect','P1','runtime-reproduced','enforce nonempty expected ID set and per-path reviewed allowlist; validate evidence binding, not shape only'))
add('tools/test-discovery-verifier.sh','ORPHAN',8,'manual-source-review','counts discovered JS but runs only run.js, smoke.js and production-verifier; discovery is not execution',('certification blocker','P1','manual-source-review','compare discovered runnable suites to actual execution ledger; classify helpers/exclusions explicitly'))
add('.github/workflows/strict-verification.yml','hidden skip',30,'manual-source-review','production truth-gate conditional on environment/secret; absent prerequisite skips required runtime evidence',('certification blocker','P1','manual-source-review','required release job must fail/NOT VERIFIED on missing secret; optional PR job must not count as runtime pass'))
findings.sort(key=lambda r:(r['file'],r['line'],r['kind'],r['origin']))
for i,r in enumerate(findings,1):r['id']=f'A37-{i:05d}'
# Small gz CSV preserves full inventory; workbook is the main browsable inventory.
fields=list(findings[0]);
with gzip.open(OUT/'inventory.csv.gz','wt',encoding='utf8',newline='') as f:
 w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(findings)
(TMP/'findings.json').write_text(json.dumps(findings,ensure_ascii=False))
with gzip.open(OUT/'evidence/ast-summary.json.gz','wt') as f:json.dump({k:{'checks':len(v.get('checks',[])),'throws':len(v.get('throws',[])),'markers':len(v.get('markers',[])),'parse_error':v.get('error')} for k,v in facts.items()},f)
with (OUT/'test-inventory.csv').open('w') as f:
 w=csv.writer(f);w.writerow(['file','raw_assert_absent','named_check_calls','throw_statements','no_recognized_oracle','direct_CI','potential_CI_graph','manual_broad_runner_graph','role'])
 for t in tests:w.writerow([t,t in narrow,len(facts[t].get('checks',[])),len(facts[t].get('throws',[])),t in noOracle,t in direct,t in reachable,t in manualReach,'helper-candidate' if helper(t) else 'suite-candidate'])
with gzip.open(OUT/'evidence/reference-graph.csv.gz','wt') as f:
 w=csv.writer(f);w.writerow(['from','to','line','kind']);w.writerows(sorted(edges))
summary={'head':(OUT/'evidence/HEAD').read_text().strip(),'text_files':len(files),'js_files':len(facts),'parse_errors':{k:v['error'] for k,v in facts.items() if 'error'in v},'test_js_cjs':len(tests),'top_level_test_js':sum(t.count('/')==1 and t.endswith('.js') for t in tests),'raw_occurrences':dict(raw),'raw_files':{k:len(v) for k,v in rawfiles.items()},'ast_markers':dict(collections.Counter(m['kind'] for v in facts.values() for m in v.get('markers',[]))),'tests_without_raw_assert':len(narrow),'tests_without_named_check_calls':len(noNamed),'tests_no_recognized_oracle_candidates':len(noOracle),'direct_ci_test_files':len(set(tests)&direct),'ci_graph_test_files':len(set(tests)&reachable),'ci_graph_unreached_candidates':len(orphan),'manual_runner_graph_test_files':len(set(tests)&manualReach),'legacy_parity_unwired_top_level':sum(t.count('/')==1 and t.endswith('.js') and t not in legacydirect for t in tests),'inventory_rows':len(findings),'classifications':dict(collections.Counter(r['classification'] for r in findings)),'review_status':dict(collections.Counter(r['confidence'] for r in findings))}
(OUT/'summary.json').write_text(json.dumps(summary,indent=2));print(json.dumps(summary,indent=2));
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
wb=Workbook();ws=wb.active;ws.title='Findings';ws.append(fields)
for r in findings:ws.append([r[k] for k in fields])
ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
widths=[14,21,62,8,25,25,10,23,62,75,95]
for i,width in enumerate(widths,1):ws.column_dimensions[get_column_letter(i)].width=width
for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17324D')
for row in ws.iter_rows(min_row=2):
 for c in row:
  if isinstance(c.value,str):c.data_type='s'
  c.alignment=Alignment(vertical='top',wrap_text=True)
notes=wb.create_sheet('Read me');
for row in [['A-37','Read-only full tracked+unignored repository scan'],['HEAD',summary['head']],['WARNING','Static candidates are not all defects. needs-review is NOT VERIFIED, not legitimate clearance.'],['ZERO-CHECK','Raw assertion absence is not runtime zero checks. Helpers and custom oracles are classified separately.'],['ORPHAN','CI graph is conditional/overapproximated. Static absence is not proof of never executed.'],['MOCK','Fixture presence is not automatically a defect; live-runtime claims require independent evidence.'],['Counts','Raw occurrences, AST constructs, and grouped inventory rows have different units.'],['Business code','No repository files modified by this audit. Existing dirty tree is preserved.']]:notes.append(row)
notes.column_dimensions['A'].width=20;notes.column_dimensions['B'].width=120
wb.save(OUT/'inventory.xlsx')
